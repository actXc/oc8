from __future__ import annotations

import sys
from pathlib import Path

import pytest

PLUGIN_ROOT = Path(__file__).resolve().parents[1]


def _evict() -> None:
    """Drop every cached `connector`/`mcp_bridge` module from sys.modules.

    Both names are generic -- after the package restructure every plugin ships
    a package called one of them (design §5.2) -- and sys.modules is keyed by
    NAME, not by path. Called SYMMETRICALLY around the import below: before,
    so a sibling plugin's cached copy cannot answer it; after, so nothing
    generic is left cached for anyone else. The "after" half is the
    load-bearing one -- `loader.import_entry_point` (Task 6's collision fix)
    only evicts modules IT ITSELF introduced, so a name left cached here makes
    a later `find_plugin`/`load_plugin` for gdrive_source or google_workspace
    silently hand back THIS plugin's code. Verified live. The module objects
    bound by the import block stay valid across the eviction.
    """
    for _stale in [
        n
        for n in sys.modules
        if n in {"connector", "mcp_bridge"} or n.startswith(("connector.", "mcp_bridge."))
    ]:
        del sys.modules[_stale]


_evict()
sys.path.insert(0, str(PLUGIN_ROOT))

from mcp_bridge import graph  # noqa: E402
from mcp_bridge.tools.powerpoint import CALL_HANDLERS, TOOLS  # noqa: E402

sys.path.remove(str(PLUGIN_ROOT))
_evict()


def test_declares_exactly_two_powerpoint_tools() -> None:
    """Deliberately just two -- see the design doc §4.2 and §8: Graph has no
    slide- or placeholder-level PowerPoint editing API, unlike Google Slides."""
    names = {t.name for t in TOOLS}
    assert names == {"powerpoint_get_text", "powerpoint_replace_file"}
    assert set(CALL_HANDLERS) == names


@pytest.mark.asyncio
async def test_get_text_extracts_slide_text(monkeypatch: pytest.MonkeyPatch) -> None:
    import io

    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Roadmap"
    buf = io.BytesIO()
    prs.save(buf)

    async def fake_get_bytes(path: str) -> bytes:
        return buf.getvalue()

    monkeypatch.setattr(graph, "get_bytes", fake_get_bytes, raising=False)
    result = await CALL_HANDLERS["powerpoint_get_text"]({"driveId": "d1", "itemId": "f1"})
    assert "Roadmap" in result
