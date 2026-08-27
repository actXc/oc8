from __future__ import annotations

import sys
from pathlib import Path

import httpx
import pytest

from oc8.knowledge.connectors.base import Attestation, ConnectorError
from oc8.oauth import http as oauth_http

PLUGIN_ROOT = Path(__file__).resolve().parents[1]


def _evict() -> None:
    """Drop every cached `connector`/`mcp_bridge` module from sys.modules.

    Both names are generic -- after the package restructure every plugin ships
    a package called one of them (design §5.2) -- and sys.modules is keyed by
    NAME, not by path. Called SYMMETRICALLY around the import below: before,
    so a sibling plugin's cached copy cannot answer it; after, so nothing
    generic is left cached for anyone else. The "after" half is the
    load-bearing one -- `loader.import_entry_point` (Task 6's collision fix)
    only evicts modules IT ITSELF introduced, so a name left cached here makes
    a later `find_plugin`/`load_plugin` for gdrive_source or google_workspace
    silently hand back THIS plugin's code. Verified live. The module objects
    bound by the import block stay valid across the eviction.
    """
    for _stale in [
        n
        for n in sys.modules
        if n in {"connector", "mcp_bridge"} or n.startswith(("connector.", "mcp_bridge."))
    ]:
        del sys.modules[_stale]


_evict()
sys.path.insert(0, str(PLUGIN_ROOT))

from connector.connector import Microsoft365FilesConnector  # noqa: E402

sys.path.remove(str(PLUGIN_ROOT))
_evict()

pytestmark = pytest.mark.asyncio


class _FakeAuth:
    """A stand-in `AuthContext`. `secret` is never called by this connector
    (Graph auth is pure OAuth) but the protocol requires it, and now that the
    plugin's own package resolves for mypy, an incomplete stub is a real
    type error rather than a silently-`Any` one."""

    async def token(self) -> str:
        return "fake-graph-token"

    async def secret(self, ref: str) -> str:
        raise AssertionError(f"microsoft365 never resolves a secret ref ({ref!r})")


def test_type_id_and_requires_oauth() -> None:
    c = Microsoft365FilesConnector()
    assert c.type_id == "microsoft365_files"
    assert c.requires_oauth == "microsoft"


async def test_validate_rejects_no_site_ids() -> None:
    c = Microsoft365FilesConnector()
    result = await c.validate({"siteIds": []}, _FakeAuth())
    assert result.ok is False


async def test_validate_rejects_missing_auth() -> None:
    c = Microsoft365FilesConnector()
    result = await c.validate({"siteIds": ["site1"]}, None)
    assert result.ok is False
    assert "azure" in (result.error or "").lower() or "microsoft" in (result.error or "").lower()


async def test_attest_listing_is_authoritative_and_enumerates_every_page() -> None:
    calls: list[str] = []

    def handler(request: httpx.Request) -> httpx.Response:
        calls.append(str(request.url))
        if "$skiptoken" not in str(request.url) and "page2" not in str(request.url):
            return httpx.Response(
                200,
                json={
                    "value": [{"id": "f1", "name": "a.txt", "file": {}}],
                    "@odata.nextLink": "https://graph.microsoft.com/v1.0/drives/d1/root/children?page2=1",
                },
            )
        return httpx.Response(200, json={"value": [{"id": "f2", "name": "b.txt", "file": {}}]})

    oauth_http.set_transport_override(httpx.MockTransport(handler))
    try:
        c = Microsoft365FilesConnector()
        listing = await c.attest_listing({"siteIds": [], "driveIds": ["d1"]}, _FakeAuth())
    finally:
        oauth_http.set_transport_override(None)

    assert listing.attestation == Attestation.AUTHORITATIVE
    assert listing.present_uris == frozenset(
        {
            "https://graph.microsoft.com/v1.0/drives/d1/items/f1",
            "https://graph.microsoft.com/v1.0/drives/d1/items/f2",
        }
    )
    assert len(calls) == 2  # both pages actually fetched


async def test_attest_listing_raises_on_api_error_rather_than_returning_empty() -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(403, json={"error": {"message": "forbidden"}})

    oauth_http.set_transport_override(httpx.MockTransport(handler))
    try:
        c = Microsoft365FilesConnector()
        with pytest.raises(ConnectorError):
            await c.attest_listing({"siteIds": [], "driveIds": ["d1"]}, _FakeAuth())
    finally:
        oauth_http.set_transport_override(None)


async def test_fetch_yields_text_files_and_skips_unsupported_types() -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        url = str(request.url)
        if url.endswith("/root/children"):
            return httpx.Response(
                200,
                json={
                    "value": [
                        {"id": "f1", "name": "a.txt", "file": {"mimeType": "text/plain"}},
                        {"id": "f2", "name": "b.png", "file": {"mimeType": "image/png"}},
                    ]
                },
            )
        if url.endswith("/items/f1/content"):
            return httpx.Response(
                200, content=b"hello world", headers={"content-type": "text/plain"}
            )
        return httpx.Response(404, json={"error": {"message": "unexpected"}})

    oauth_http.set_transport_override(httpx.MockTransport(handler))
    try:
        c = Microsoft365FilesConnector()
        docs = [doc async for doc in c.fetch({"driveIds": ["d1"]}, None, _FakeAuth())]
    finally:
        oauth_http.set_transport_override(None)

    assert len(docs) == 1
    assert docs[0].title == "a.txt"
    assert docs[0].content == "hello world"
    assert docs[0].source_uri == "https://graph.microsoft.com/v1.0/drives/d1/items/f1"


async def test_fetch_skips_a_file_already_seen_by_content_hash() -> None:
    import hashlib

    def handler(request: httpx.Request) -> httpx.Response:
        url = str(request.url)
        if url.endswith("/root/children"):
            return httpx.Response(
                200,
                json={"value": [{"id": "f1", "name": "a.txt", "file": {"mimeType": "text/plain"}}]},
            )
        return httpx.Response(200, content=b"same content", headers={"content-type": "text/plain"})

    oauth_http.set_transport_override(httpx.MockTransport(handler))
    try:
        c = Microsoft365FilesConnector()
        digest = hashlib.sha256(b"same content").hexdigest()
        docs = [
            doc async for doc in c.fetch({"driveIds": ["d1"]}, {"hashes": [digest]}, _FakeAuth())
        ]
    finally:
        oauth_http.set_transport_override(None)
    assert docs == []


async def test_fetch_extracts_text_from_a_word_document(tmp_path: Path) -> None:
    import docx

    doc = docx.Document()
    doc.add_paragraph("First paragraph.")
    doc.add_paragraph("Second paragraph.")
    docx_path = tmp_path / "doc.docx"
    doc.save(str(docx_path))
    docx_bytes = docx_path.read_bytes()

    def handler(request: httpx.Request) -> httpx.Response:
        url = str(request.url)
        if url.endswith("/root/children"):
            return httpx.Response(
                200,
                json={
                    "value": [
                        {
                            "id": "f1",
                            "name": "doc.docx",
                            "file": {
                                "mimeType": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"  # noqa: E501
                            },
                        }
                    ]
                },
            )
        return httpx.Response(200, content=docx_bytes)

    oauth_http.set_transport_override(httpx.MockTransport(handler))
    try:
        c = Microsoft365FilesConnector()
        docs = [doc async for doc in c.fetch({"driveIds": ["d1"]}, None, _FakeAuth())]
    finally:
        oauth_http.set_transport_override(None)

    assert len(docs) == 1
    assert "First paragraph." in docs[0].content
    assert "Second paragraph." in docs[0].content


async def test_fetch_extracts_text_from_an_excel_workbook(tmp_path: Path) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Revenue"
    ws["B1"] = 1000
    xlsx_path = tmp_path / "book.xlsx"
    wb.save(str(xlsx_path))
    xlsx_bytes = xlsx_path.read_bytes()

    def handler(request: httpx.Request) -> httpx.Response:
        url = str(request.url)
        if url.endswith("/root/children"):
            return httpx.Response(
                200,
                json={
                    "value": [
                        {
                            "id": "f1",
                            "name": "book.xlsx",
                            "file": {
                                "mimeType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"  # noqa: E501
                            },
                        }
                    ]
                },
            )
        return httpx.Response(200, content=xlsx_bytes)

    oauth_http.set_transport_override(httpx.MockTransport(handler))
    try:
        c = Microsoft365FilesConnector()
        docs = [doc async for doc in c.fetch({"driveIds": ["d1"]}, None, _FakeAuth())]
    finally:
        oauth_http.set_transport_override(None)

    assert len(docs) == 1
    assert "Revenue" in docs[0].content
    assert "1000" in docs[0].content


async def test_fetch_follows_the_302_graph_answers_a_content_download_with() -> None:
    """`GET /drives/{id}/items/{id}/content` does not return the bytes.

    It returns 302 + `Location` pointing at a short-lived pre-authenticated
    download URL on a *different* host. Every mock in this file used to answer
    that call with a direct 200, which asserted the API behaves the way the
    code assumed -- and against the real Graph the first ingestible file
    answered 302, `_api_error` turned it into "HTTP 302", and the entire sync
    aborted having ingested nothing (whole-branch review, C3).

    The bearer token must NOT travel to the redirect target: that URL already
    carries its own credential, and the target is a SharePoint/blob host, not
    Graph. httpx drops `Authorization` on a cross-origin hop; this pins it.
    """
    hops: list[tuple[str, str | None]] = []

    def handler(request: httpx.Request) -> httpx.Response:
        url = str(request.url)
        hops.append((url, request.headers.get("authorization")))
        if url.endswith("/root/children"):
            return httpx.Response(
                200,
                json={"value": [{"id": "f1", "name": "a.txt", "file": {"mimeType": "text/plain"}}]},
            )
        if url.endswith("/items/f1/content"):
            return httpx.Response(
                302,
                headers={"Location": "https://contoso.sharepoint.com/download?t=preauth"},
            )
        return httpx.Response(200, content=b"hello from sharepoint")

    oauth_http.set_transport_override(httpx.MockTransport(handler))
    try:
        c = Microsoft365FilesConnector()
        docs = [doc async for doc in c.fetch({"driveIds": ["d1"]}, None, _FakeAuth())]
    finally:
        oauth_http.set_transport_override(None)

    assert len(docs) == 1
    assert docs[0].content == "hello from sharepoint"
    followed = [hop for hop in hops if "sharepoint.com" in hop[0]]
    assert followed, "the redirect was never followed"
    assert followed[0][1] is None, "the Graph bearer token leaked to the redirect target"


async def test_fetch_skips_an_unreadable_office_file_instead_of_aborting_the_sync() -> None:
    """One corrupt .docx used to take the whole sync down with it.

    The connector's own stated policy is skip-and-count; a password-protected
    or truncated Office file is exactly the case that policy is for, and it is
    far likelier to be hit now that downloads actually reach the parser at all.
    The good file after it must still arrive.
    """
    docx_mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

    def handler(request: httpx.Request) -> httpx.Response:
        url = str(request.url)
        if url.endswith("/root/children"):
            return httpx.Response(
                200,
                json={
                    "value": [
                        {"id": "bad", "name": "broken.docx", "file": {"mimeType": docx_mime}},
                        {"id": "good", "name": "a.txt", "file": {"mimeType": "text/plain"}},
                    ]
                },
            )
        if url.endswith("/items/bad/content"):
            return httpx.Response(200, content=b"this is not a zip archive at all")
        return httpx.Response(200, content=b"readable text")

    oauth_http.set_transport_override(httpx.MockTransport(handler))
    try:
        c = Microsoft365FilesConnector()
        docs = [doc async for doc in c.fetch({"driveIds": ["d1"]}, None, _FakeAuth())]
    finally:
        oauth_http.set_transport_override(None)

    assert [d.title for d in docs] == ["a.txt"]


async def test_fetch_extracts_text_from_a_powerpoint_deck(tmp_path: Path) -> None:
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Quarterly Results"
    pptx_path = tmp_path / "deck.pptx"
    prs.save(str(pptx_path))
    pptx_bytes = pptx_path.read_bytes()

    def handler(request: httpx.Request) -> httpx.Response:
        url = str(request.url)
        if url.endswith("/root/children"):
            return httpx.Response(
                200,
                json={
                    "value": [
                        {
                            "id": "f1",
                            "name": "deck.pptx",
                            "file": {
                                "mimeType": "application/vnd.openxmlformats-officedocument.presentationml.presentation"  # noqa: E501
                            },
                        }
                    ]
                },
            )
        return httpx.Response(200, content=pptx_bytes)

    oauth_http.set_transport_override(httpx.MockTransport(handler))
    try:
        c = Microsoft365FilesConnector()
        docs = [doc async for doc in c.fetch({"driveIds": ["d1"]}, None, _FakeAuth())]
    finally:
        oauth_http.set_transport_override(None)

    assert len(docs) == 1
    assert "Quarterly Results" in docs[0].content
