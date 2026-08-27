"""PowerPoint tools -- deliberately just two. Microsoft Graph has no
slide-level or placeholder-level PowerPoint editing API (unlike the Excel
Workbook API, and unlike Google Slides -- see the design doc §4.2/§8 for the
stated asymmetry). `powerpoint_get_text` downloads and parses locally with
python-pptx, using the same extraction logic the RAG connector's own
`_extract_pptx_text` (Task 6) already uses -- duplicated here rather than
imported, since the connector and this MCP bridge are separate Python
packages with no shared dependency between them; `powerpoint_replace_file` is
a whole-file re-upload, the only "edit" Graph actually offers for this
format."""

from __future__ import annotations

import io
from collections.abc import Awaitable, Callable
from typing import Any

from mcp.types import Tool

from .. import graph


def _extract_pptx_text(raw: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(raw))
    lines: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text:
                lines.append(shape.text_frame.text)
    return "\n".join(lines)


async def powerpoint_get_text(args: dict[str, Any]) -> Any:
    raw = await graph.get_bytes(
        f"/drives/{args['driveId']}/items/{args['itemId']}/content"
    )
    return _extract_pptx_text(raw)


async def powerpoint_replace_file(args: dict[str, Any]) -> Any:
    content = args["content"]
    raw = content.encode() if isinstance(content, str) else content
    return await graph.put_bytes(
        f"/drives/{args['driveId']}/items/{args['itemId']}/content", raw
    )


TOOLS: list[Tool] = [
    Tool(
        name="powerpoint_get_text",
        description="Extract all text from a PowerPoint deck's slides (no slide-level structure).",
        inputSchema={
            "type": "object",
            "properties": {"driveId": {"type": "string"}, "itemId": {"type": "string"}},
            "required": ["driveId", "itemId"],
        },
    ),
    Tool(
        name="powerpoint_replace_file",
        description="Replace an entire PowerPoint file's binary content (no targeted slide edit is available).",
        inputSchema={
            "type": "object",
            "properties": {
                "driveId": {"type": "string"},
                "itemId": {"type": "string"},
                "content": {"type": "string"},
            },
            "required": ["driveId", "itemId", "content"],
        },
    ),
]

CALL_HANDLERS: dict[str, Callable[[dict[str, Any]], Awaitable[Any]]] = {
    "powerpoint_get_text": powerpoint_get_text,
    "powerpoint_replace_file": powerpoint_replace_file,
}
